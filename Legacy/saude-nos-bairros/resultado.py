#!/usr/bin/env python3
"""Resultado: confirmações por especialidade por dia -> resultado.xlsx.

Uso: uv run resultado.py [saida.xlsx]
Confirmação = 'Houve interação=Sim' em QUALQUER relatório de Relatorios/.
Dia = Data do agendamento em Dados/. Especialidade = Categorias split ';'
(Plano) / Categoria/Meta (RNM). Multi-especialidade conta em cada uma.
Telefone compartilhado com Sim: só a 1ª linha conta; demais = indefinidas
(fora do numerador e do denominador).
"""
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import numbers

from recaptura import BASE, DADOS, REL, fmt_categoria, fmt_date, norm_phone, read_csv


def main(out=None):
    out = Path(out or BASE / "resultado.xlsx")
    gerado = datetime.now().strftime("%d/%m/%Y %H:%M")
    sim = set()
    STATUS_FALHA = {"A mensagem falhou", "O número de telefone não possui WhatsApp"}
    disparos = []  # (nome, enviadas, entregues, lidas, enviado, interagidas, falhas)
    for f in sorted(REL.glob("*.csv")):
        rows = read_csv(f)
        enviadas = len(rows)
        entregues = lidas = enviado = falhas = interagidas = 0
        for r in rows:
            status = (r.get("Status da mensagem", "") or "").strip()
            if (r.get("Houve interação", "") or "").strip().lower() == "sim":
                interagidas += 1
                if p := norm_phone(r.get("Contato/Telefone", "")):
                    sim.add(p)
            if status == "Lida":
                lidas += 1
            elif status == "Entregue":
                entregues += 1
            elif status == "Enviado":
                enviado += 1
            elif status in STATUS_FALHA:
                falhas += 1
        assert enviadas == entregues + lidas + enviado + falhas, f"assert total {f.name}"
        disparos.append([f.name, enviadas, entregues, lidas, enviado, interagidas, falhas])

    # titularidade: 1ª linha do telefone vale; resto = indefinido
    seen, titulares, indefinidos, invalidos = set(), [], 0, 0
    for f in sorted(DADOS.glob("*.csv")):
        for r in read_csv(f):
            cols = r.keys()
            tel = r.get("Telefone padronizado") if "Telefone padronizado" in cols else r.get("Telefone", "")
            p = norm_phone(tel)
            if len(p) not in (10, 11):
                invalidos += 1
                continue
            if p in seen:
                if p in sim:
                    indefinidos += 1  # interação não atribuível
                    continue
                # Não: cada paciente segue precisando de contato
            else:
                seen.add(p)
            titulares.append({"phone": p,
                              "data": fmt_date(r.get("Data", "")),
                              "specs": [s for s in dict.fromkeys(fmt_categoria(
                                  r.get("Categoria/Meta") or r.get("Categorias") or "").split("; ")) if s],
                              "conf": p in sim})

    conv = defaultdict(int)  # (spec, dia) -> convocados atribuíveis
    conf = defaultdict(int)  # (spec, dia) -> confirmados
    for t in titulares:
        for s in t["specs"]:
            conv[(s, t["data"])] += 1
            if t["conf"]:
                conf[(s, t["data"])] += 1
    specs = sorted({s for s, _ in conv})
    dias = sorted({d for _, d in conv}, key=lambda d: datetime.strptime(d, "%d/%m/%Y"))

    wb = Workbook()
    w1 = wb.active
    w1.title = "Confirmados"
    w1.append(["Especialidade", *dias, "Total"])
    for s in specs:
        w1.append([s, *(conf[(s, d)] for d in dias), sum(conf[(s, d)] for d in dias)])
    w1.append(["TOTAL", *(sum(conf[(s, d)] for s in specs) for d in dias),
               sum(conf.values())])
    w1.freeze_panes = "B2"
    w1.auto_filter.ref = w1.dimensions

    w2 = wb.create_sheet("Taxa")
    w2.append(["Especialidade", *dias, "Total"])
    for s in specs:
        row = [s]
        for d in dias:
            # ponytail: sem convocados -> vazio (não 0%, que significa outra coisa)
            row.append(conf[(s, d)] / conv[(s, d)] if conv[(s, d)] else None)
        tot_c = sum(conv[(s, d)] for d in dias)
        row.append(sum(conf[(s, d)] for d in dias) / tot_c if tot_c else 0)
        w2.append(row)
    tot_all = sum(conv.values())
    w2.append(["TOTAL", *(sum(conf[(s, d)] for s in specs) / sum(conv[(s, d)] for s in specs)
                           if sum(conv[(s, d)] for s in specs) else None for d in dias),
               sum(conf.values()) / tot_all if tot_all else 0])
    for row in w2.iter_rows(min_row=2, min_col=2):
        for c in row:
            c.number_format = numbers.FORMAT_PERCENTAGE_00
    w2.freeze_panes = "B2"
    w2.column_dimensions["A"].width = 28

    w4 = wb.create_sheet("Disparos")
    w4.append(["Relatório (disparo)", "Enviadas", "Entregues", "Lidas",
               "Enviado s/ confirmação", "Interagidas", "Falhas"])
    tot = [0] * 6
    for d in disparos:
        w4.append(d)
        for i in range(6):
            tot[i] += d[i + 1]
    w4.append(["TOTAL", *tot])
    for c in "BCDEFG":
        w4.column_dimensions[c].width = 12
    w4.column_dimensions["A"].width = 80

    w3 = wb.create_sheet("Resumo")
    for k, v in [("Gerado em", gerado),
                 ("Mensagens enviadas (total, todos disparos)", tot[0]),
                 ("Falhas (total)", tot[-1]),
                 ("Telefones Sim únicos (confirmados)", len(sim)),
                 ("Titulares atribuíveis", len(titulares)),
                 ("Linhas indefinidas (compartilhadas c/ Sim)", indefinidos),
                 ("Telefones inválidos (ignorados)", invalidos),
                 ("Confirmações (titulares Sim)", sum(1 for t in titulares if t["conf"])),
                 ("Taxa global", f"{sum(1 for t in titulares if t['conf']) / len(titulares):.1%}" if titulares else "—")]:
        w3.append([k, v])
    w3.column_dimensions["A"].width = 46
    wb.save(out)

    # validação (matrizes explodem multi-especialidade: 1 titular = N células)
    assert sum(conf.values()) == sum(len(t["specs"]) for t in titulares if t["conf"])
    assert sum(1 for t in titulares if t["conf"]) <= len(sim), "mais titulares que Sim únicos"
    assert all(conf[k] <= conv[k] for k in conf)
    print(f"gerado={gerado} specs={len(specs)} dias={len(dias)} "
          f"titulares={len(titulares)} conf={sum(conf.values())} indef={indefinidos} -> {out.name}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else None)
