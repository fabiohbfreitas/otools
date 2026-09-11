"""Biblioteca da recaptação: leitura, distribuição e escrita (usada por envios.py)."""
import glob, os, re, sys
from collections import Counter, defaultdict

import openpyxl

INPUT_HDR = ["Paciente", "Telefone", "Data", "Horário", "Unidade", "Especialidade", "Situação", "Observação"]


def split_phones(raw):
    if not raw or not str(raw).strip():
        return []
    return [p.strip() for p in str(raw).split(",") if p.strip()]


def pick_main(phones):
    # ponytail: 11 dígitos = móvel 9 dígitos; upgrade se regra de DDD mudar
    if not phones:
        return None, []
    for p in phones:
        if len(re.sub(r"\D", "", p)) == 11:
            rest = list(phones)
            rest.remove(p)
            return p, rest
    return phones[0], phones[1:]


def agenda_str(slot):
    hh, mm = slot.split(":")
    return f"Agenda{hh}h" if mm == "00" else f"Agenda{hh}h{mm}"


def assign_slots(items, slots, min_n):
    # ponytail: espalha nos horários necessários; mínimo min_n por slot, salvo resto final
    n = len(items)
    if n == 0:
        return []
    k = min(len(slots), n // min_n) if n >= min_n else 1
    k = max(k, 1)
    base, rem = divmod(n, k)
    out, j = [], 0
    for s in range(k):
        for _ in range(base + (1 if s < rem else 0)):
            out.append((items[j], slots[s]))
            j += 1
    return out


def local_link(cfg, polo, esp):
    p = cfg["polos"][polo]
    ov = (p.get("especialidades") or {}).get(esp) or {}
    return ov.get("endereco", p["endereco"]), ov.get("link", p["link"])


def load_inputs(indir):
    pats = [indir] if os.path.isfile(indir) else sorted(glob.glob(os.path.join(indir, "*.xlsx")))
    if not pats:
        sys.exit(f"nenhum xlsx em {indir}")
    out = []
    for f in pats:
        wb = openpyxl.load_workbook(f, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        hdr = [str(c).strip() if c else "" for c in rows[0][:8]]
        if hdr[:6] != INPUT_HDR[:6]:
            sys.exit(f"{f}: cabeçalho inesperado: {rows[0][:8]}")
        for r in rows[1:]:
            r = list(r) + [None] * 8  # Situação/Observação opcionais
            if not r[0] or not str(r[0]).strip():
                continue
            out.append({"paciente": str(r[0]).strip(), "tel_raw": str(r[1]).strip() if r[1] else "",
                        "data_recap": str(r[2]).strip() if r[2] else "", "horario": str(r[3]).strip() if r[3] else "",
                        "polo": str(r[4]).strip(), "esp": str(r[5]).strip() if r[5] else "",
                        "situacao": str(r[6]).strip() if r[6] else "", "obs": str(r[7]).strip() if r[7] else ""})
    return out


def plan_quotas(total, ndates, meta=0):
    # ponytail: sem meta = divisão igual; com meta = exato por dia, sobra p/ excedentes
    if not meta:
        base, rem = divmod(total, ndates)
        return [base + (1 if i < rem else 0) for i in range(ndates)]
    quotas, left = [], total
    for _ in range(ndates):
        q = min(meta, left)
        quotas.append(q)
        left -= q
    return quotas


def distribute(patients, dates, polos_por_dia, slots, min_horario, metas=None):
    metas = metas or {}
    by_polo = defaultdict(list)
    for p in patients:
        by_polo[p["polo"]].append(p)
    buckets = {(d, polo): [] for d, pls in zip(dates, polos_por_dia) for polo in pls}
    exced = []
    for polo, lst in by_polo.items():
        ds = [d for d, pls in zip(dates, polos_por_dia) if polo in pls]
        if not ds:
            sys.exit(f"polo {polo} sem data destino")
        i = 0
        for d, n in zip(ds, plan_quotas(len(lst), len(ds), metas.get(polo, 0))):
            buckets[(d, polo)] = assign_slots(lst[i:i + n], slots, min_horario)
            i += n
        exced.extend(lst[i:])
    return buckets, exced


def validate_config(cfg):
    cfg_min = cfg.get("min_por_horario", 4)
    metas = cfg.get("metas_por_polo", {})
    for polo, m in metas.items():
        if polo not in cfg["polos"]:
            sys.exit(f"metas_por_polo: polo sem config: {polo}")
        if not isinstance(m, int) or m <= 0:
            sys.exit(f"metas_por_polo: {polo} deve ser inteiro > 0")
    for hora in cfg["slots"]:
        try:
            hh, mm = hora.split(":")
            int(hh)
            int(mm)
        except (ValueError, AttributeError):
            sys.exit(f"slots: hora inválida '{hora}'")
    return cfg_min, metas


def check_metas(buckets, dates, polos_por_dia, metas, min_n=0):
    warns = []
    for d, pls in zip(dates, polos_por_dia):
        for polo in pls:
            bloco = buckets.get((d, polo), [])
            if metas.get(polo, 0) and len(bloco) < metas[polo]:
                warns.append(f"AVISO: {polo} em {d}: {len(bloco)} < meta {metas[polo]}")
            if min_n:
                for h, k in sorted(Counter(h for _, h in bloco).items()):
                    if k < min_n:
                        warns.append(f"AVISO: {polo} em {d} {h}: grupo com {k} < mínimo {min_n}")
    return warns


def write_excedentes(path, pacientes):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Excedentes"
    ws.append(INPUT_HDR)
    for p in pacientes:
        ws.append([p["paciente"], p["tel_raw"] or None, p["data_recap"] or None, p["horario"] or None,
                   p["polo"], p["esp"] or None, p["situacao"] or None, p["obs"] or None])
    wb.save(path)


def selfcheck():
    assert pick_main(["(61) 3485-9002", "(61) 99587-0517", "(61) 99132-0964"])[0] == "(61) 99587-0517"
    assert pick_main(["(61) 3315-2425", "(61) 9938-2494", "(61) 99636-9724"])[0] == "(61) 99636-9724"
    assert pick_main(["(61) 9169-6717", "(61) 99169-6717"])[0] == "(61) 99169-6717"
    assert pick_main([]) == (None, [])
    assert split_phones(None) == []
    assert agenda_str("09:00") == "Agenda09h" and agenda_str("13:00") == "Agenda13h"
    assert agenda_str("07:30") == "Agenda07h30" and agenda_str("15:15") == "Agenda15h15"
    assert [s for _, s in assign_slots(list(range(10)), ["00", "01"], 4)] == ["00"] * 5 + ["01"] * 5
    assert assign_slots(list(range(3)), ["a", "b"], 4) == [(0, "a"), (1, "a"), (2, "a")]
    assert [s for _, s in assign_slots(list(range(100)), ["a", "b", "c"], 5)] == ["a"] * 34 + ["b"] * 33 + ["c"] * 33
    b = {("2026-09-01", "Gama"): [1] * 5, ("2026-09-01", "Sobradinho"): [1] * 40}
    assert check_metas(b, ["2026-09-01"], [["Gama", "Sobradinho"]], {"Gama": 70}) == ["AVISO: Gama em 2026-09-01: 5 < meta 70"]
    assert check_metas(b, ["2026-09-01"], [["Gama", "Sobradinho"]], {}) == []
    assert plan_quotas(242, 2, 70) == [70, 70]
    assert plan_quotas(441, 2, 154) == [154, 154]
    assert plan_quotas(200, 3, 70) == [70, 70, 60]
    assert plan_quotas(10, 3) == [4, 3, 3]
    print("selfcheck ok")
