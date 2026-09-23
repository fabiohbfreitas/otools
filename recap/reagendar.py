"""Reagendar: troca os horários de um Pacientes+Envios já gerado, preservando o resto."""
import argparse, datetime, json, os, sys
from collections import Counter
from copy import copy

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from recap import agenda_str, assign_slots, esp_base


def check_slot(s):
    try:
        hh, mm = s.split(":")
        int(hh)
        int(mm)
    except ValueError:
        sys.exit(f"horário inválido: '{s}'")
    return s


def parse_grupos(spec):
    grupos = []
    for g in spec.split(";"):
        n, _, sl = g.partition(":")
        slots = [check_slot(s.strip()) for s in sl.split(",") if s.strip()]
        if not n.strip().isdigit() or not slots:
            sys.exit(f"grupo inválido: '{g}' (use N:hh:mm,hh:mm;...)")
        grupos.append((int(n.strip()), slots))
    return grupos


def parse_cotas(spec):
    cotas = {}
    for c in spec.split(","):
        s, _, n = c.partition("=")
        if not n.strip().isdigit() or not s.strip():
            sys.exit(f"cota inválida: '{c}' (use hh:mm=N,...)")
        cotas[check_slot(s.strip())] = int(n.strip())
    return cotas


def distribute_slots(n, slots, min_n, grupos=None, cotas=None):
    if grupos is not None:
        if sum(k for k, _ in grupos) != n:
            sys.exit(f"grupos somam {sum(k for k, _ in grupos)}, arquivo tem {n}")
        out, i = [None] * n, 0
        for k, gsl in grupos:
            for j, s in assign_slots(list(range(k)), gsl, min_n):
                out[i + j] = s
            i += k
        return out
    if cotas:
        for s in cotas:
            if s not in slots:
                sys.exit(f"cota fora da lista de horários: {s}")
        if sum(cotas.values()) > n:
            sys.exit(f"cotas somam {sum(cotas.values())}, arquivo tem {n}")
        livres = [s for s in slots if s not in cotas]
        out, i = [], 0
        for s, q in cotas.items():
            out += [s] * q
            i += q
        if i < n:
            if not livres:
                sys.exit(f"sobram {n - i} sem horário livre")
            out += [s for _, s in assign_slots(list(range(n - i)), livres, min_n)]
        return out
    return [s for _, s in assign_slots(list(range(n)), slots, min_n)]


def iso_of(val):
    if isinstance(val, str):
        d, m, y = val.split("/")
        return f"{y}-{m}-{d}"
    return val.strftime("%Y-%m-%d")


def set_hora(cell, slot):
    hh, mm = slot.split(":")
    cell.value = f"{hh}:{mm}" if isinstance(cell.value, str) else datetime.time(int(hh), int(mm))


def set_data(cell, iso):
    dt = datetime.date.fromisoformat(iso)
    if isinstance(cell.value, str):
        cell.value = dt.strftime("%d/%m/%Y")
    else:
        cell.value = dt


def col_idx(hdr, name, fname):
    if name not in hdr:
        sys.exit(f"{fname}: coluna '{name}' não encontrada")
    return hdr.index(name)


def output_names(input_path, sufixo=""):
    base, _ = os.path.splitext(input_path)
    if sufixo.strip():
        base += f" {sufixo.strip()}"
    return f"{base} - Pacientes.xlsx", f"{base} - Envios.xlsx"


def save_sheet(src_ws, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = src_ws.title
    for row in src_ws.iter_rows():
        for c in row:
            nc = ws.cell(row=c.row, column=c.column, value=c.value)
            if c.has_style:
                nc._style = copy(c._style)
    for col, dim in src_ws.column_dimensions.items():
        ws.column_dimensions[col].width = dim.width
    wb.save(path)


def selfcheck():
    assert parse_grupos("20:07:30,08:15;17:13:00,14:00") == [(20, ["07:30", "08:15"]), (17, ["13:00", "14:00"])]
    assert parse_cotas("07:30=10,08:15=5") == {"07:30": 10, "08:15": 5}
    assert distribute_slots(10, ["07:30", "08:15"], 4) == ["07:30"] * 5 + ["08:15"] * 5
    assert distribute_slots(37, ["a", "b"], 4, grupos=[(20, ["a"]), (17, ["b"])]) == ["a"] * 20 + ["b"] * 17
    assert distribute_slots(14, ["07:30", "08:15"], 4, cotas={"07:30": 10}) == ["07:30"] * 10 + ["08:15"] * 4
    assert output_names("a/b.xlsx", "") == ("a/b - Pacientes.xlsx", "a/b - Envios.xlsx")
    assert output_names("a/b.xlsx", "v2") == ("a/b v2 - Pacientes.xlsx", "a/b v2 - Envios.xlsx")
    print("selfcheck ok")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=False, help="xlsx com abas Pacientes+Envios")
    ap.add_argument("--horarios", default="", help='"07:30,08:15,..." (divisão igual)')
    ap.add_argument("--grupos", default="", help='"20:07:30,08:15;17:13:00,14:00"')
    ap.add_argument("--cotas", default="", help='"07:30=10,..." (exige --horarios)')
    ap.add_argument("--nova-data", default="", help="AAAA-MM-DD (opcional)")
    ap.add_argument("--sufixo", default="", help='ex: "v2" (nome dos arquivos)')
    ap.add_argument("--config", default="config.json")
    ap.add_argument("--selfcheck", action="store_true")
    a = ap.parse_args()
    if a.selfcheck:
        return selfcheck()
    if not a.input:
        ap.error("--input é obrigatório")
    if sum(bool(x) for x in [a.horarios, a.grupos]) != 1:
        ap.error("use exatamente um: --horarios ou --grupos")
    if a.cotas and not a.horarios:
        ap.error("--cotas exige --horarios")
    if a.nova_data:
        datetime.date.fromisoformat(a.nova_data)
    cfg = json.load(open(a.config, encoding="utf-8"))
    min_n = cfg.get("min_por_horario", 4)
    slots = [check_slot(s.strip()) for s in a.horarios.split(",") if s.strip()] if a.horarios else []
    grupos = parse_grupos(a.grupos) if a.grupos else None
    cotas = parse_cotas(a.cotas) if a.cotas else None

    wb = openpyxl.load_workbook(a.input)
    if "Pacientes" not in wb.sheetnames or "Envios" not in wb.sheetnames:
        sys.exit("esperava abas Pacientes+Envios")
    wp, we = wb["Pacientes"], wb["Envios"]
    hp = [c.value for c in wp[1]]
    he = [c.value for c in we[1]]
    i_nome = col_idx(hp, "Nome", "Pacientes")
    i_hora = col_idx(hp, "Hora", "Pacientes")
    i_data_p = col_idx(hp, "Data", "Pacientes")
    i_esp = col_idx(hp, "Especialidade", "Pacientes")
    i_loc = col_idx(hp, "Local", "Pacientes")
    i_hor = col_idx(he, "[horario]", "Envios")
    i_eti = col_idx(he, "Etiquetas", "Envios")
    i_dat = col_idx(he, "[data]", "Envios")
    rows_p = [r for r in wp.iter_rows(min_row=2) if r[i_nome].value]
    rows_e = [r for r in we.iter_rows(min_row=2) if r[0].value]
    if len(rows_p) != len(rows_e):
        sys.exit(f"Pacientes ({len(rows_p)}) e Envios ({len(rows_e)}) divergem")
    novo = distribute_slots(len(rows_p), slots, min_n, grupos, cotas)
    for rp, re, slot in zip(rows_p, rows_e, novo):
        iso = a.nova_data or iso_of(rp[i_data_p].value)
        polo = str(rp[i_loc].value).strip()
        base = esp_base(str(rp[i_esp].value) if rp[i_esp].value else "")
        etiq = cfg["etiqueta_envios"].format(data_iso=iso, polo=polo, esp=base, agenda=agenda_str(slot))
        set_hora(rp[i_hora], slot)
        set_hora(re[i_hor], slot)
        re[i_eti].value = etiq
        if a.nova_data:
            set_data(rp[i_data_p], iso)
            set_data(re[i_dat], iso)
    f_pac, f_env = output_names(a.input, a.sufixo)
    save_sheet(wp, f_pac)
    save_sheet(we, f_env)
    for s, k in sorted(Counter(novo).items()):
        print(f"  {s}: {k}")
    print(f"{len(novo)} pacientes -> {f_pac} | {f_env}")


if __name__ == "__main__":
    main()
